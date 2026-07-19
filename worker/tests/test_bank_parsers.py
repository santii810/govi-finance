from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from worker.preview import analyze_file
from worker.parsers.revolut import parse_file


def _write_revolut_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Tipo,Producto,Fecha de inicio,Fecha de finalización,Descripción,Importe,Comisión,Divisa,State,Saldo"
        ]
    )
    sheet.append(
        [
            "Pago con tarjeta,Actual,2026-06-01 10:00:00,2026-06-01 12:00:00,Cafetería,-4.50,0.00,EUR,COMPLETADO,100.00"
        ]
    )
    sheet.append(
        [
            "Recargas,Actual,2026-06-02 09:00:00,2026-06-02 09:01:00,Recarga,200.00,0.00,EUR,COMPLETADO,300.00"
        ]
    )
    workbook.save(path)
    workbook.close()


def _write_myinvestor_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Fecha de operación", "Fecha de valor", "Concepto", "Importe", "Divisa"])
    sheet.append(["45809", "45809", "RECARGA MI", "200", "EUR"])
    sheet.append(["45810", "45810", "VANGUARD US 500", "-24.54", "EUR"])
    workbook.save(path)
    workbook.close()


def test_analyze_revolut_xlsx(tmp_path: Path):
    sample = tmp_path / "revolut.csv"
    _write_revolut_xlsx(sample)

    preview, movements = analyze_file(sample)

    assert preview.account_id == "revolut-santi"
    assert preview.banco == "Revolut"
    assert preview.ambiguedades
    assert preview.total_movimientos == 2
    assert preview.gastos_count == 1
    assert preview.ingresos_count == 1
    assert preview.fecha_desde == date(2026, 6, 1)
    assert preview.fecha_hasta == date(2026, 6, 2)
    assert movements[0].importe == Decimal("-4.50")
    assert "Cafetería" in movements[0].concepto
    assert len(movements[0].idempotency_key) == 32


def test_analyze_revolut_comun(tmp_path: Path):
    sample = tmp_path / "revolut.csv"
    _write_revolut_xlsx(sample)

    preview, movements = analyze_file(sample, account_id="revolut-comun")

    assert preview.account_id == "revolut-comun"
    assert preview.persona == "Común"
    assert preview.tipo_cuenta == "conjunta"
    assert movements[0].persona == "Común"


def test_revolut_prefers_fecha_inicio_over_fecha_fin(tmp_path: Path):
    sample = tmp_path / "revolut.csv"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Tipo,Producto,Fecha de inicio,Fecha de finalización,Descripción,Importe,Comisión,Divisa,State,Saldo"
        ]
    )
    sheet.append(
        [
            "Reembolso de tarjeta,Actual,2025-12-30 11:00:00,2026-01-01 13:00:00,Bricomart,8.33,0.00,EUR,COMPLETADO,100.00"
        ]
    )
    workbook.save(sample)
    workbook.close()

    movements = parse_file(sample)

    assert len(movements) == 1
    assert movements[0].fecha == date(2025, 12, 30)


def test_analyze_myinvestor_xlsx(tmp_path: Path):
    sample = tmp_path / "myinvestor.csv"
    _write_myinvestor_xlsx(sample)

    preview, movements = analyze_file(sample)

    assert preview.account_id == "myinvestor-santi"
    assert preview.banco == "MyInvestor"
    assert preview.total_movimientos == 2
    assert preview.gastos_count == 1
    assert preview.ingresos_count == 1
    assert movements[0].importe == Decimal("200")
    assert movements[1].importe == Decimal("-24.54")


def test_analyze_myinvestor_semicolon_csv(tmp_path: Path):
    """MyInvestor exporta CSV texto con separador ';' e importes con coma decimal."""
    sample = tmp_path / "Movimientos Mi Cuenta MyInvestor.csv"
    sample.write_text(
        "Fecha de operación;Fecha de valor;Concepto;Importe;Divisa\r\n"
        "28/04/2026;29/04/2026;VANGUARD US 500 STOCK INDEX EU;-24,8;EUR\r\n"
        "07/04/2026;07/04/2026;RECAGA MYINVESTOR;110;EUR\r\n",
        encoding="utf-8",
    )

    preview, movements = analyze_file(sample, account_id="myinvestor-sandra")

    assert preview.account_id == "myinvestor-sandra"
    assert preview.banco == "MyInvestor"
    assert preview.persona == "Sandra"
    assert preview.total_movimientos == 2
    assert preview.gastos_count == 1
    assert preview.ingresos_count == 1
    assert movements[0].importe == Decimal("-24.8")
    assert movements[1].importe == Decimal("110")


def test_myinvestor_excel_and_csv_share_idempotency_keys(tmp_path: Path):
    """Mismo movimiento en Excel (formato viejo) y CSV ';' (formato nuevo) → misma key."""
    xlsx = tmp_path / "old.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Fecha de operación", "Fecha de valor", "Concepto", "Importe", "Divisa"])
    sheet.append(["2026-04-28 00:00:00", "2026-04-29 00:00:00", "VANGUARD US 500", "-24.8", "EUR"])
    sheet.append(["2026-04-07 00:00:00", "2026-04-07 00:00:00", "RECAGA MYINVESTOR", "110", "EUR"])
    workbook.save(xlsx)
    workbook.close()

    csv_path = tmp_path / "new.csv"
    csv_path.write_text(
        "Fecha de operación;Fecha de valor;Concepto;Importe;Divisa\r\n"
        "28/04/2026;29/04/2026;VANGUARD US 500;-24,8;EUR\r\n"
        "07/04/2026;07/04/2026;RECAGA MYINVESTOR;110;EUR\r\n",
        encoding="utf-8",
    )

    _, from_xlsx = analyze_file(xlsx, account_id="myinvestor-santi")
    _, from_csv = analyze_file(csv_path, account_id="myinvestor-santi")

    assert [m.idempotency_key for m in from_xlsx] == [m.idempotency_key for m in from_csv]
    assert from_xlsx[0].idempotency_key == "2026-04-28|2026-04-29|VANGUARD US 500|-24.8"
    assert from_xlsx[1].idempotency_key == "2026-04-07|2026-04-07|RECAGA MYINVESTOR|110"


@pytest.mark.parametrize(
    "sample_path,account_id,min_movements",
    [
        (
            Path("samples/revolut-santi/account-statement_2021-12-01_2026-06-13_es-es_29cca1 (2).csv"),
            "revolut-santi",
            450,
        ),
        (
            Path("samples/myinvestor-santi/Movimientos Mi Cuenta MyInvestor.csv"),
            "myinvestor-santi",
            120,
        ),
    ],
)
def test_real_samples_if_present(sample_path: Path, account_id: str, min_movements: int):
    if not sample_path.is_file():
        pytest.skip("muestra local no disponible")

    preview, movements = analyze_file(sample_path)

    assert preview.account_id == account_id
    assert preview.total_movimientos >= min_movements
    assert len(movements) == preview.total_movimientos
