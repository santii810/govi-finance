from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from worker.excel_import.config import SHEETS


def load_sheet_rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Hoja «{sheet_name}» no encontrada en {path.name}")
        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        return []

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        if row is None or not any(cell not in (None, "") for cell in row):
            continue
        record: dict[str, Any] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            value = row[index] if index < len(row) else None
            record[header] = value
        if record:
            records.append(record)
    return records


def load_workbook_sheets(path: Path, entities: list[str]) -> dict[str, list[dict[str, Any]]]:
    return {entity: load_sheet_rows(path, SHEETS[entity]) for entity in entities}
