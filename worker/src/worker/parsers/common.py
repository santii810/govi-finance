from __future__ import annotations

import hashlib
import io
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


def is_excel_file(path: Path) -> bool:
    with path.open("rb") as handle:
        return handle.read(2) == b"PK"


def read_excel_matrix(path: Path) -> list[list[str]]:
    with path.open("rb") as handle:
        payload = io.BytesIO(handle.read())
    workbook = load_workbook(payload, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell) for cell in row])
        return rows
    finally:
        workbook.close()


def split_csv_line(line: str) -> list[str]:
    parts: list[str] = []
    current: list[str] = []
    in_quotes = False
    for char in line:
        if char == '"':
            in_quotes = not in_quotes
            continue
        if char == "," and not in_quotes:
            parts.append("".join(current).strip())
            current = []
            continue
        current.append(char)
    parts.append("".join(current).strip())
    return parts


def normalize_header(value: str) -> str:
    normalized = value.strip().lower()
    normalized = (
        normalized.replace("ã³", "ó")
        .replace("Ã³", "ó")
        .replace("Ã", "í")
        .replace("ã", "í")
    )
    replacements = {
        "fecha de operación": "fecha_operacion",
        "fecha de operacion": "fecha_operacion",
        "fecha de valor": "fecha_valor",
        "fecha de finalización": "fecha_fin",
        "fecha de finalizacion": "fecha_fin",
        "fecha de inicio": "fecha_inicio",
        "descripción": "descripcion",
        "descripcion": "descripcion",
        "comisión": "comision",
        "comision": "comision",
        "concepto": "concepto",
        "importe": "importe",
        "divisa": "divisa",
        "tipo": "tipo",
        "producto": "producto",
        "state": "state",
        "saldo": "saldo",
    }
    return replacements.get(normalized, normalized)


def parse_decimal(value: str) -> Decimal:
    cleaned = value.strip().replace("€", "").replace("EUR", "").strip()
    cleaned = cleaned.replace("\u00a0", "").replace(" ", "")
    if cleaned.count(",") == 1 and cleaned.count(".") == 0:
        cleaned = cleaned.replace(",", ".")
    elif cleaned.count(".") > 1 and cleaned.count(",") == 1:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    try:
        return Decimal(cleaned or "0")
    except InvalidOperation as exc:
        raise ValueError(f"Importe no reconocido: {value!r}") from exc


def parse_date_flexible(value: str) -> date:
    value = value.strip()
    if not value:
        raise ValueError("Fecha vacía")
    if value.isdigit() or _looks_like_excel_serial(value):
        return excel_serial_to_date(value)
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d.%m.%Y",
    ):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Fecha no reconocida: {value!r}")


def excel_serial_to_date(value: str) -> date:
    serial = int(float(value.strip()))
    return date(1899, 12, 30) + timedelta(days=serial)


def _looks_like_excel_serial(value: str) -> bool:
    try:
        number = float(value.strip())
    except ValueError:
        return False
    return number > 10_000


def hash_fallback(concepto: str, fecha: date, importe: Decimal) -> str:
    payload = f"{fecha.isoformat()}|{importe}|{concepto}"
    return hashlib.sha256(payload.encode()).hexdigest()[:32]
